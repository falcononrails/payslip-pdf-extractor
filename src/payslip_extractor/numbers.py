from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

SUPPORTED_NUMBER_FILE_SUFFIXES = {".csv", ".txt", ".xlsx", ".xlsm"}


class NumberFileError(ValueError):
    """Raised when the number input file cannot be read."""


def _is_searchable_identifier(text: str) -> bool:
    return bool(text) and any(c.isdigit() for c in text) and text.replace("_", "").isalnum()


def read_numbers(path: Path) -> list[str]:
    """Read unique searchable identifiers from Excel, CSV, or plain-text-like files."""
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in {".xlsx", ".xlsm"}:
        values = _iter_workbook_values(path)
    else:
        values = _iter_text_values(path)

    return _collect_numbers(values, "number file")


def read_numbers_from_text(text: str) -> list[str]:
    """Read unique searchable identifiers from pasted plain text."""
    return _collect_numbers(_iter_text_tokens(text), "pasted identifiers")


def _collect_numbers(values: Iterable[object], source_label: str) -> list[str]:
    seen: set[str] = set()
    numbers: list[str] = []
    for value in values:
        candidate = normalize_number_cell(value)
        if candidate and candidate not in seen:
            seen.add(candidate)
            numbers.append(candidate)

    if not numbers:
        raise NumberFileError(f"No searchable identifiers were found in {source_label}.")

    return numbers


def normalize_number_cell(value: object) -> str | None:
    """Convert one spreadsheet cell value into a searchable identifier."""
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        text = format(value, "f").rstrip("0").rstrip(".")
        return text if _is_searchable_identifier(text) else None

    if isinstance(value, Decimal):
        if value == value.to_integral_value():
            return str(value.to_integral_value())
        text = format(value.normalize(), "f").rstrip("0").rstrip(".")
        return text if _is_searchable_identifier(text) else None

    text = str(value).strip().removeprefix("'").strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text if _is_searchable_identifier(text) else None


def _iter_text_values(path: Path) -> Iterable[object]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = path.read_text(encoding="cp1252")
        except UnicodeDecodeError as exc:
            raise NumberFileError(f"Could not read text from number file '{path}'.") from exc
    except OSError as exc:
        raise NumberFileError(f"Could not read number file '{path}': {exc}") from exc

    yield from _iter_text_tokens(text)


def _iter_text_tokens(text: str) -> Iterable[object]:
    for token in re.findall(r"[A-Za-z0-9_'.]+", text):
        yield token


def _iter_workbook_values(path: Path) -> Iterable[object]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl exposes several parser-specific exceptions.
        raise NumberFileError(f"Could not read Excel file '{path}': {exc}") from exc

    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                yield from row
    finally:
        workbook.close()
