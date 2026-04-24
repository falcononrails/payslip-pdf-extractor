from __future__ import annotations

from openpyxl import Workbook

from payslip_extractor.numbers import normalize_number_cell, read_numbers


def test_normalize_number_cell_keeps_digit_strings_and_integral_numbers() -> None:
    assert normalize_number_cell("001234") == "001234"
    assert normalize_number_cell("'987") == "987"
    assert normalize_number_cell(1234) == "1234"
    assert normalize_number_cell(1234.0) == "1234"
    assert normalize_number_cell("1234.0") == "1234"


def test_normalize_number_cell_keeps_alphanumeric_identifiers() -> None:
    assert normalize_number_cell("456A") == "456A"
    assert normalize_number_cell("A12345") == "A12345"
    assert normalize_number_cell("CNSS456") == "CNSS456"
    assert normalize_number_cell("AB123CD") == "AB123CD"
    assert normalize_number_cell("001ABC") == "001ABC"


def test_normalize_number_cell_ignores_headers_and_non_digit_values() -> None:
    assert normalize_number_cell("matricule") is None
    assert normalize_number_cell("CNSS 123") is None
    assert normalize_number_cell("Name") is None
    assert normalize_number_cell(True) is None
    assert normalize_number_cell(None) is None


def test_read_numbers_from_csv_deduplicates_in_order(tmp_path) -> None:
    csv_file = tmp_path / "numbers.csv"
    csv_file.write_text("matricule,001\n123,001\n456,\n", encoding="utf-8")

    assert read_numbers(csv_file) == ["001", "123", "456"]


def test_read_numbers_from_csv_includes_alphanumeric(tmp_path) -> None:
    csv_file = tmp_path / "numbers.csv"
    csv_file.write_text("456A\n123\nA789\nName\n", encoding="utf-8")

    assert read_numbers(csv_file) == ["456A", "123", "A789"]


def test_read_numbers_from_xlsx_all_sheets(tmp_path) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first.append(["matricule", "123"])
    first.append([456, None])
    second = workbook.create_sheet("Second")
    second.append(["789", "123"])

    xlsx_file = tmp_path / "numbers.xlsx"
    workbook.save(xlsx_file)
    workbook.close()

    assert read_numbers(xlsx_file) == ["123", "456", "789"]


def test_read_numbers_from_xlsx_with_alphanumeric(tmp_path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["456A", "B789", "header", "123"])

    xlsx_file = tmp_path / "numbers.xlsx"
    workbook.save(xlsx_file)
    workbook.close()

    assert read_numbers(xlsx_file) == ["456A", "B789", "123"]